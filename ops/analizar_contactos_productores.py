"""Analiza un padrón de contactos sin importar ni exponer datos personales.

El informe por stdout contiene sólo totales. Las filas que requieren intervención
se guardan en un CSV privado indicado por ``--revision``.
"""

from __future__ import annotations

import argparse
import csv
import os
import re
import unicodedata
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path

import psycopg
from openpyxl import load_workbook


@dataclass(frozen=True)
class Contacto:
    fila: int
    suc: str
    nombre: str
    telefonos: tuple[str, ...]
    observacion: str


def texto(valor: object) -> str:
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def nombre_exacto(valor: object) -> str:
    """Normalización deliberadamente conservadora para coincidencia automática."""
    return " ".join(unicodedata.normalize("NFC", texto(valor)).upper().split())


def suc_normalizada(valor: object) -> str:
    bruto = texto(valor).upper()
    return re.sub(r"[^A-Z0-9]", "", bruto)


def normalizar_telefono(valor: object) -> tuple[set[str], bool]:
    """Devuelve E.164 argentinos inequívocos y si quedó contenido sin resolver."""
    bruto = texto(valor)
    if not bruto:
        return set(), False

    encontrados: set[str] = set()
    ambiguo = False
    partes = re.split(r"[\n;,/]+", bruto)
    for parte in partes:
        digitos = re.sub(r"\D", "", parte)
        if not digitos:
            continue
        digitos = digitos.removeprefix("00")
        if digitos.startswith("0") and len(digitos) == 11:
            digitos = digitos[1:]

        if len(digitos) == 13 and digitos.startswith("549"):
            normalizado = digitos
        elif len(digitos) == 12 and digitos.startswith("54"):
            normalizado = "549" + digitos[2:]
        elif len(digitos) == 11 and digitos.startswith("9"):
            normalizado = "54" + digitos
        elif len(digitos) == 10:
            normalizado = "549" + digitos
        else:
            ambiguo = True
            continue
        encontrados.add(normalizado)
    return encontrados, ambiguo


def leer_contactos(ruta: Path) -> tuple[list[Contacto], Counter[str]]:
    libro = load_workbook(ruta, read_only=True, data_only=True)
    hoja = libro["CONTACTOS"]
    contactos: list[Contacto] = []
    metricas: Counter[str] = Counter()
    for numero_fila, valores in enumerate(
        hoja.iter_rows(min_row=2, max_col=4, values_only=True), start=2
    ):
        suc, nombre, primario, alternativo = valores
        telefonos: set[str] = set()
        ambiguo = False
        for valor in (primario, alternativo):
            hallados, dudoso = normalizar_telefono(valor)
            telefonos.update(hallados)
            ambiguo |= dudoso
        if ambiguo:
            metricas["filas_telefono_ambiguo"] += 1
        metricas[f"filas_con_{min(len(telefonos), 2)}_telefonos"] += 1
        contactos.append(
            Contacto(
                fila=numero_fila,
                suc=suc_normalizada(suc),
                nombre=nombre_exacto(nombre),
                telefonos=tuple(sorted(telefonos)),
                observacion=texto(alternativo),
            )
        )
    return contactos, metricas


def cargar_productores(dsn: str) -> list[tuple[str, str, str]]:
    with psycopg.connect(dsn) as conexion, conexion.cursor() as cursor:
        cursor.execute(
            """SELECT DISTINCT clienterazonsocial, clientecuit, nro_delegacion
               FROM recepcion.descarga
               WHERE clienterazonsocial IS NOT NULL
                 AND btrim(clienterazonsocial) <> ''
                 AND clientecuit IS NOT NULL
                 AND btrim(clientecuit) <> ''"""
        )
        return [
            (nombre_exacto(nombre), texto(cuit), suc_normalizada(delegacion))
            for nombre, cuit, delegacion in cursor.fetchall()
        ]


def analizar(
    contactos: list[Contacto], productores: list[tuple[str, str, str]]
) -> tuple[
    Counter[str],
    list[dict[str, str]],
    list[dict[str, str]],
    dict[str, set[str]],
]:
    por_nombre: dict[str, set[str]] = defaultdict(set)
    por_nombre_suc: dict[tuple[str, str], set[str]] = defaultdict(set)
    delegaciones_por_suc_excel: dict[str, set[str]] = defaultdict(set)
    for nombre, cuit, delegacion in productores:
        por_nombre[nombre].add(cuit)
        por_nombre_suc[(nombre, delegacion)].add(cuit)

    metricas: Counter[str] = Counter()
    revision: list[dict[str, str]] = []
    asignacion_telefono: dict[str, set[str]] = defaultdict(set)
    origenes_telefono: dict[str, list[Contacto]] = defaultdict(list)
    for contacto in contactos:
        candidatos = por_nombre.get(contacto.nombre, set())
        candidatos_suc = por_nombre_suc.get((contacto.nombre, contacto.suc), set())
        motivo = ""
        cuit = ""
        if len(candidatos) == 1:
            cuit = next(iter(candidatos))
            metricas["coincidencia_exacta_unica"] += 1
            delegaciones_por_suc_excel[contacto.suc].update(
                delegacion
                for nombre, cuit_db, delegacion in productores
                if nombre == contacto.nombre and cuit_db == cuit
            )
        elif len(candidatos) > 1 and len(candidatos_suc) == 1:
            # Se mide, pero no se autoacepta hasta validar que SUC sea delegación.
            metricas["desambiguable_solo_si_suc_es_delegacion"] += 1
            motivo = "nombre_varios_cuit_suc_candidata"
        elif len(candidatos) > 1:
            metricas["nombre_ambiguo"] += 1
            motivo = "nombre_varios_cuit"
        else:
            metricas["sin_coincidencia_exacta"] += 1
            motivo = "nombre_sin_coincidencia"

        if cuit and contacto.telefonos:
            for telefono in contacto.telefonos:
                asignacion_telefono[telefono].add(cuit)
                origenes_telefono[telefono].append(contacto)
        else:
            if not contacto.telefonos:
                motivo = f"{motivo};sin_telefono_valido".strip(";")
            revision.append(
                {
                    "fila_excel": str(contacto.fila),
                    "suc": contacto.suc,
                    "productor": contacto.nombre,
                    "telefonos": "|".join(contacto.telefonos),
                    "observacion": contacto.observacion,
                    "motivo": motivo,
                }
            )

    conflictos = {t: cuits for t, cuits in asignacion_telefono.items() if len(cuits) > 1}
    filas_conflictivas: set[int] = set()
    for telefono in conflictos:
        for contacto in origenes_telefono[telefono]:
            if contacto.fila in filas_conflictivas:
                continue
            filas_conflictivas.add(contacto.fila)
            revision.append(
                {
                    "fila_excel": str(contacto.fila),
                    "suc": contacto.suc,
                    "productor": contacto.nombre,
                    "telefonos": "|".join(contacto.telefonos),
                    "observacion": contacto.observacion,
                    "motivo": "telefono_asociado_a_varios_cuit",
                }
            )
    aprobados = [
        {"telefono": telefono, "clientecuit": next(iter(cuits))}
        for telefono, cuits in sorted(asignacion_telefono.items())
        if len(cuits) == 1
    ]
    metricas["telefonos_unicos_asignables"] = sum(
        1 for cuits in asignacion_telefono.values() if len(cuits) == 1
    )
    metricas["telefonos_con_conflicto_de_cuit"] = len(conflictos)
    metricas["filas_revision"] = len(revision)
    return metricas, revision, aprobados, delegaciones_por_suc_excel


def guardar_revision(ruta: Path, filas: list[dict[str, str]]) -> None:
    ruta.parent.mkdir(mode=0o700, parents=True, exist_ok=True)
    fd = os.open(ruta, os.O_WRONLY | os.O_CREAT | os.O_TRUNC, 0o600)
    with os.fdopen(fd, "w", newline="", encoding="utf-8") as archivo:
        campos = [
            "fila_excel",
            "suc",
            "productor",
            "telefonos",
            "observacion",
            "motivo",
        ]
        escritor = csv.DictWriter(archivo, fieldnames=campos)
        escritor.writeheader()
        escritor.writerows(filas)


def guardar_aprobados(ruta: Path, filas: list[dict[str, str]]) -> None:
    ruta.parent.mkdir(mode=0o700, parents=True, exist_ok=True)
    fd = os.open(ruta, os.O_WRONLY | os.O_CREAT | os.O_TRUNC, 0o600)
    with os.fdopen(fd, "w", newline="", encoding="utf-8") as archivo:
        escritor = csv.DictWriter(archivo, fieldnames=["telefono", "clientecuit"])
        escritor.writeheader()
        escritor.writerows(filas)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--excel", type=Path, required=True)
    parser.add_argument("--dsn", default=os.environ.get("RECEPCION_DSN_SUITE"))
    parser.add_argument("--revision", type=Path, required=True)
    parser.add_argument("--aprobados", type=Path, required=True)
    args = parser.parse_args()
    if not args.dsn:
        parser.error("falta --dsn o RECEPCION_DSN_SUITE")

    contactos, telefonos = leer_contactos(args.excel)
    productores = cargar_productores(args.dsn)
    resultado, revision, aprobados, correlacion_suc = analizar(contactos, productores)
    guardar_revision(args.revision, revision)
    guardar_aprobados(args.aprobados, aprobados)

    print(f"filas_excel={len(contactos)}")
    print(f"filas_db_distintas={len(productores)}")
    for clave in sorted(telefonos):
        print(f"{clave}={telefonos[clave]}")
    for clave in sorted(resultado):
        print(f"{clave}={resultado[clave]}")
    print(f"suc_excel_distintas={len(correlacion_suc)}")
    print(
        "suc_con_correspondencia_univoca_delegacion="
        + str(sum(1 for delegaciones in correlacion_suc.values() if len(delegaciones) == 1))
    )
    print(f"revision_privada={args.revision}")
    print(f"aprobados_privados={args.aprobados}")


if __name__ == "__main__":
    main()
